from __future__ import annotations
import json
from io import BytesIO
from uuid import UUID, uuid4
from fastapi import APIRouter, Depends, HTTPException
from fastapi.responses import FileResponse as FastAPIFileResponse
from sqlalchemy import select, insert, update
from sqlalchemy.ext.asyncio import AsyncSession
from pathlib import Path
from app.core.auth import CurrentUser, require_roles
from app.core.config import settings
from app.db.session import get_db
from app.models.models import metadata
from app.schemas.requests import ReportCreateRequest
from app.schemas.schemas import ReportJobResponse

router=APIRouter(prefix="/reports",tags=["reports"])
ALLOWED_FORMATS={"xls","xlsx","pdf","json","png"}
ALLOWED_TABLES={"universities","programs","interactions","tasks","documents","activities","users"}

def _query_table(db, table_name, filters):
    table=metadata.tables.get(table_name)
    if table is None: raise HTTPException(422,f"Unsupported report table: {table_name}")
    stmt=select(table)
    for key,val in filters.items():
        if key not in table.c: raise HTTPException(422,f"Unknown filter field: {key}")
        stmt=stmt.where(table.c[key]==val)
    return table,stmt

async def _generate(job, db):
    fmt=job['format']; filters=job['filters'] or {}; table_name=filters.get('table','universities')
    if table_name not in ALLOWED_TABLES: raise HTTPException(422,'Unsupported report table')
    table,stmt=_query_table(db,table_name,{k:v for k,v in filters.items() if k!='table'})
    rows=[dict(x) for x in (await db.execute(stmt.limit(5000))).mappings().all()]
    cols=job['selected_columns'] or [c.name for c in table.columns]
    cols=[c for c in cols if c in table.c]
    if not cols: raise HTTPException(422,'selected_columns contains no valid columns')
    matrix=[[r.get(c) for c in cols] for r in rows]
    outdir=Path(settings.storage_path)/'reports'; outdir.mkdir(parents=True,exist_ok=True)
    jid=job['id']; path=outdir/f'{jid}.{fmt}'
    if fmt=='json':
        path.write_text(json.dumps([{c:r.get(c) for c in cols} for r in rows],default=str,ensure_ascii=False,indent=2),encoding='utf-8')
        ctype='application/json'
    elif fmt=='xlsx':
        from openpyxl import Workbook
        wb=Workbook(); ws=wb.active; ws.title=table_name; ws.append(cols)
        for row in matrix: ws.append(row)
        wb.save(path); ctype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    elif fmt=='xls':
        import xlwt
        wb=xlwt.Workbook(); ws=wb.add_sheet(table_name[:31]); ws.write(0,0,cols[0])
        for j,c in enumerate(cols): ws.write(0,j,c)
        for i,row in enumerate(matrix,1):
            for j,val in enumerate(row): ws.write(i,j,str(val) if val is not None else '')
        wb.save(str(path)); ctype='application/vnd.ms-excel'
    elif fmt=='pdf':
        from reportlab.platypus import SimpleDocTemplate, Table, TableStyle
        from reportlab.lib import colors
        doc=SimpleDocTemplate(str(path)); data=[cols]+[[str(v) if v is not None else '' for v in row] for row in matrix]
        tbl=Table(data,repeatRows=1); tbl.setStyle(TableStyle([('GRID',(0,0),(-1,-1),0.5,colors.grey),('BACKGROUND',(0,0),(-1,0),colors.lightgrey)])); doc.build([tbl]); ctype='application/pdf'
    else:
        from PIL import Image, ImageDraw
        text=table_name+'\n'+' | '.join(cols)+'\n'+'\n'.join(' | '.join(str(v) for v in row) for row in matrix[:50])
        img=Image.new('RGB',(1600,max(300,40*(len(text.splitlines())+1))), 'white'); ImageDraw.Draw(img).multiline_text((20,20),text,fill='black'); img.save(path); ctype='image/png'
    ftable=metadata.tables['files']; fid=uuid4(); key=f'reports/{path.name}'
    raw=path.read_bytes(); import hashlib
    await db.execute(insert(ftable).values(id=fid,storage_key=key,original_name=path.name,content_type=ctype,size_bytes=len(raw),checksum_sha256=hashlib.sha256(raw).hexdigest(),uploaded_by=job['requested_by']))
    await db.execute(update(metadata.tables['report_jobs']).where(metadata.tables['report_jobs'].c.id==jid).values(status='completed',result_file_id=fid,completed_at=__import__('sqlalchemy').func.now()))
    await db.commit(); return fid

@router.post("",response_model=ReportJobResponse,status_code=202)
async def create_report(payload:ReportCreateRequest,user:CurrentUser=Depends(require_roles("admin","manager","user")),db:AsyncSession=Depends(get_db)):
    if payload.format not in ALLOWED_FORMATS: raise HTTPException(422,'format must be xls, xlsx, pdf, json or png')
    t=metadata.tables['report_jobs']
    job=(await db.execute(insert(t).values(requested_by=user.id,format=payload.format,status='queued',filters=payload.filters,selected_columns=payload.selected_columns).returning(t))).mappings().one()
    await db.commit()
    try: await _generate(job,db)
    except Exception as exc:
        await db.rollback(); await db.execute(update(t).where(t.c.id==job['id']).values(status='failed',error_message=str(exc),completed_at=__import__('sqlalchemy').func.now())); await db.commit()
    return dict((await db.execute(select(t).where(t.c.id==job['id']))).mappings().one())

@router.get("",response_model=list[ReportJobResponse])
async def list_reports(user:CurrentUser=Depends(require_roles("admin","manager","user")),db:AsyncSession=Depends(get_db)):
    t=metadata.tables['report_jobs']; rows=(await db.execute(select(t).where(t.c.requested_by==user.id).order_by(t.c.created_at.desc()).limit(100))).mappings().all()
    return [dict(x) for x in rows]

@router.get("/{job_id}",response_model=ReportJobResponse)
async def get_report(job_id:UUID,user:CurrentUser=Depends(require_roles("admin","manager","user")),db:AsyncSession=Depends(get_db)):
    t=metadata.tables['report_jobs']; row=(await db.execute(select(t).where(t.c.id==job_id))).mappings().first()
    if not row: raise HTTPException(404,'Report job not found')
    if row['requested_by']!=user.id and 'admin' not in user.roles: raise HTTPException(403,'Not allowed')
    return dict(row)
